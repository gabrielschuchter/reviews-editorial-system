"""Exportação DOCX simples, nativa para Google Docs e verificável."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path
from typing import Any


class DocxExportError(RuntimeError):
    pass


def _imports():
    try:
        from docx import Document
        from docx.enum.section import WD_SECTION
        from docx.enum.style import WD_STYLE_TYPE
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:
        raise DocxExportError(
            "Exportação DOCX requer o extra opcional: pip install -e .[documents]"
        ) from exc
    return {
        "Document": Document,
        "WD_ALIGN_PARAGRAPH": WD_ALIGN_PARAGRAPH,
        "OxmlElement": OxmlElement,
        "qn": qn,
        "Inches": Inches,
        "Pt": Pt,
        "RGBColor": RGBColor,
    }


def _set_run_font(run: Any, name: str, size: Any, color: Any, bold: bool | None = None) -> None:
    run.font.name = name
    run.font.size = size
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    qn = _imports()["qn"]
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)


def _configure_styles(document: Any, api: dict[str, Any]) -> None:
    Pt = api["Pt"]
    RGBColor = api["RGBColor"]
    black = RGBColor(0, 0, 0)
    gray = RGBColor(67, 67, 67)
    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = black
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    settings = {
        "Heading 1": (20, black, 20, 6),
        "Heading 2": (16, black, 18, 6),
        "Heading 3": (14, gray, 16, 4),
    }
    for name, (size, color, before, after) in settings.items():
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = False
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(11)
        style.font.color.rgb = black
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.15


INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def _add_inline(paragraph: Any, text: str, api: dict[str, Any]) -> None:
    Pt = api["Pt"]
    RGBColor = api["RGBColor"]
    for part in filter(None, INLINE_RE.split(text)):
        run = paragraph.add_run()
        if part.startswith("**") and part.endswith("**"):
            run.text = part[2:-2]
            run.bold = True
        elif part.startswith("*") and part.endswith("*"):
            run.text = part[1:-1]
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run.text = part[1:-1]
        else:
            run.text = part
        _set_run_font(run, "Arial", Pt(11), RGBColor(0, 0, 0), run.bold)


def _parse_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _is_table_separator(line: str) -> bool:
    cells = _parse_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def _column_widths(rows: list[list[str]], total: int = 9360) -> list[int]:
    count = max(len(row) for row in rows)
    weights = [max(8, max((len(row[index]) if index < len(row) else 0) for row in rows)) for index in range(count)]
    raw = [max(900, round(total * weight / sum(weights))) for weight in weights]
    difference = total - sum(raw)
    raw[-1] += difference
    if raw[-1] < 900:
        deficit = 900 - raw[-1]
        raw[-1] = 900
        raw[raw.index(max(raw[:-1]))] -= deficit
    return raw


def _set_cell_width(cell: Any, width: int, api: dict[str, Any]) -> None:
    OxmlElement, qn = api["OxmlElement"], api["qn"]
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")
    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", 80), ("start", 120), ("bottom", 80), ("end", 120)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _add_table(document: Any, rows: list[list[str]], api: dict[str, Any]) -> None:
    OxmlElement, qn = api["OxmlElement"], api["qn"]
    widths = _column_widths(rows)
    table = document.add_table(rows=len(rows), cols=len(widths))
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "DADCE0")
        borders.append(node)
    tbl_pr.append(borders)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row_index, values in enumerate(rows):
        for column_index, cell in enumerate(table.rows[row_index].cells):
            _set_cell_width(cell, widths[column_index], api)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = api["Pt"](0)
            paragraph.paragraph_format.space_after = api["Pt"](4)
            paragraph.paragraph_format.line_spacing = 1.15
            _add_inline(paragraph, values[column_index] if column_index < len(values) else "", api)
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True
    document.add_paragraph().paragraph_format.space_after = api["Pt"](0)


def _strip_frontmatter(lines: list[str]) -> list[str]:
    if lines and lines[0].strip() == "---":
        for index, line in enumerate(lines[1:], start=1):
            if line.strip() == "---":
                return lines[index + 1 :]
    return lines


def export_markdown_to_docx(markdown_path: str | Path, output_path: str | Path) -> dict[str, Any]:
    api = _imports()
    source = Path(markdown_path)
    lines = _strip_frontmatter(source.read_text(encoding="utf-8-sig").splitlines())
    document = api["Document"]()
    section = document.sections[0]
    section.page_width = api["Inches"](8.5)
    section.page_height = api["Inches"](11)
    section.top_margin = section.right_margin = section.bottom_margin = section.left_margin = api["Inches"](1)
    section.header_distance = section.footer_distance = api["Inches"](0.492)
    _configure_styles(document, api)

    title_text: str | None = None
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        if (
            "|" in line
            and index + 1 < len(lines)
            and _is_table_separator(lines[index + 1])
        ):
            table_rows = [_parse_table_row(line)]
            index += 2
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                table_rows.append(_parse_table_row(lines[index]))
                index += 1
            _add_table(document, table_rows, api)
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2).strip()
            if level == 1 and title_text is None:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_before = api["Pt"](0)
                paragraph.paragraph_format.space_after = api["Pt"](3)
                paragraph.paragraph_format.keep_with_next = True
                run = paragraph.add_run(text)
                _set_run_font(run, "Arial", api["Pt"](26), api["RGBColor"](0, 0, 0), False)
                title_text = text
            else:
                paragraph = document.add_paragraph(style=f"Heading {min(level, 3)}")
                _add_inline(paragraph, text, api)
            index += 1
            continue
        bullet = re.match(r"^\s*[-*+]\s+(.+)$", line)
        numbered = re.match(r"^\s*\d+[.)]\s+(.+)$", line)
        if bullet or numbered:
            paragraph = document.add_paragraph(style="List Bullet" if bullet else "List Number")
            _add_inline(paragraph, (bullet or numbered).group(1), api)
            index += 1
            continue
        if not line.strip():
            index += 1
            continue
        paragraph_lines = [line.strip()]
        index += 1
        while index < len(lines) and lines[index].strip() and not re.match(
            r"^(#{1,4})\s+|^\s*[-*+]\s+|^\s*\d+[.)]\s+", lines[index]
        ):
            if "|" in lines[index] and index + 1 < len(lines) and _is_table_separator(lines[index + 1]):
                break
            paragraph_lines.append(lines[index].strip())
            index += 1
        paragraph = document.add_paragraph()
        _add_inline(paragraph, " ".join(paragraph_lines), api)

    document.core_properties.title = title_text or source.stem
    document.core_properties.subject = "Candidata editorial do Reviews"
    document.core_properties.comments = "Exportado pelo Reviews Editorial System; revisão humana obrigatória."
    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
    audit = audit_google_docs_docx(target)
    if not audit["passed"]:
        raise DocxExportError("DOCX falhou na auditoria estrutural: " + "; ".join(audit["issues"]))
    return audit


def audit_google_docs_docx(path: str | Path) -> dict[str, Any]:
    issues: list[str] = []
    with zipfile.ZipFile(path) as package:
        document_xml = package.read("word/document.xml").decode("utf-8", errors="replace")
        styles_xml = package.read("word/styles.xml").decode("utf-8", errors="replace")
        first_paragraph = document_xml.split("</w:p>", 1)[0]
        if "w:pBdr" in first_paragraph:
            issues.append("o bloco de título contém borda de parágrafo")
        if 'w:pStyle w:val="Title"' in first_paragraph:
            issues.append("o título usa o estilo Word Title")
        if "Arial" not in styles_xml:
            issues.append("Arial não foi codificada nos estilos")
        if "w:pgSz" not in document_xml or "w:pgMar" not in document_xml:
            issues.append("geometria de página ausente")
    return {"passed": not issues, "issues": issues, "preset": "google_docs_default"}
